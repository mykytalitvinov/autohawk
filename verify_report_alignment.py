from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

from src.scanner import AutohawkScanner
from src.best_of_scan_report import write_best_of_scan_report


ROOT = Path(__file__).resolve().parent


def _load_config() -> dict:
    return json.loads((ROOT / "config.json").read_text(encoding="utf-8-sig"))


def _excel_links(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    link_col = headers.index("Link") + 1
    links: list[str] = []
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row=row, column=link_col).value
        if value:
            links.append(str(value))
    return links


def _text_links(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    return [match.group(1) for match in re.finditer(r"(?m)^\s*Link:\s*(https?://\S+)", text)]


def main() -> int:
    config = _load_config()
    scanner = AutohawkScanner(config)
    scanner._refresh_learned_market()
    scanner._update_excel()
    write_best_of_scan_report(config, scanner.scan_count, scanner._last_export_listings)

    excel_path = ROOT / config.get("output_file", "output/deals.xlsx")
    text_path = ROOT / config.get("best_of_scan_report_file", "output/best_of_scan.txt")

    excel = _excel_links(excel_path)
    text = _text_links(text_path)
    compared = min(len(excel), len(text), int(config.get("best_of_scan_rows", 10)))
    mismatches = [
        (idx + 1, excel[idx], text[idx])
        for idx in range(compared)
        if excel[idx] != text[idx]
    ]

    print(f"Excel rows: {len(excel)}")
    print(f"Best-of-scan rows: {len(text)}")
    print(f"Compared top rows: {compared}")
    print(f"Mismatches: {len(mismatches)}")
    for idx, excel_link, text_link in mismatches[:10]:
        print(f"#{idx}: Excel={excel_link} TXT={text_link}")

    if mismatches:
        return 1
    print("OK: deals.xlsx and best_of_scan.txt use the same top order.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
